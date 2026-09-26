import vk_api
from vk_api.bot_longpoll import VkBotLongPoll, VkBotEventType
from vk_api.keyboard import VkKeyboard, VkKeyboardColor
import sqlite3
import time
import os
from datetime import datetime

# ======== ТОКЕН И ID ГРУППЫ ========
TOKEN = "vk1.a.VGjAHYDuV0YBY2hKZsNNoldS1o9Ce-83n6ROFJj0M5GJPVyqufFIdvLH-oJbXBzUCEFR0VXIMT-IqW6zfdRhxAAolq3_pjmh6h8wAgpHQpZanvAFpKblN8-d_lUB8N-jj5xnGJHNage0wy4y7totMCRgFu1TCjXdMX7FqcjWlO4_xBgIVfy9sPuwiCeKm54cjdP1xU9iWBhL4oBv3oVnMA" 
GROUP_ID = 237002976 

# ======== ПОДКЛЮЧЕНИЕ К ВК ========
vk_session = vk_api.VkApi(token=TOKEN)
longpoll = VkBotLongPoll(vk_session, GROUP_ID)
vk = vk_session.get_api()

users = {}

# ======== РАБОТА С EXCEL ========
try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    EXCEL_AVAILABLE = True
except ImportError:
    print("⚠️ openpyxl не установлена! Установите: pip3 install openpyxl")
    EXCEL_AVAILABLE = False

EXCEL_FILE = "заявки.xlsx"

def create_excel():
    wb = Workbook()
    ws = wb.active
    ws.title = "Заявки"
    
    header_fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    header_font = Font(bold=True, size=12, color="FFFFFF")
    center_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), 
                         top=Side(style='thin'), bottom=Side(style='thin'))
    
    headers = ["№", "Дата", "ФИО", "Соавтор", "Учебное заведение", 
               "Образование", "Секция", "Название статьи", "Руководитель", 
               "Файл статьи", "Справка антиплагиат", "Презентация"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
    
    ws.column_dimensions['A'].width = 5
    ws.column_dimensions['B'].width = 18
    ws.column_dimensions['C'].width = 25
    ws.column_dimensions['D'].width = 25
    ws.column_dimensions['E'].width = 30
    ws.column_dimensions['F'].width = 15
    ws.column_dimensions['G'].width = 60
    ws.column_dimensions['H'].width = 40
    ws.column_dimensions['I'].width = 25
    ws.column_dimensions['J'].width = 40
    ws.column_dimensions['K'].width = 40
    ws.column_dimensions['L'].width = 40
    
    wb.save(EXCEL_FILE)

def add_to_excel(data):
    if not EXCEL_AVAILABLE:
        return
    
    if not os.path.exists(EXCEL_FILE):
        create_excel()
    
    try:
        wb = load_workbook(EXCEL_FILE)
        ws = wb.active
        row = ws.max_row + 1
        
        ws.cell(row=row, column=1, value=row - 1)
        ws.cell(row=row, column=2, value=datetime.now().strftime("%d.%m.%Y %H:%M"))
        ws.cell(row=row, column=3, value=data.get("full_name", ""))
        ws.cell(row=row, column=4, value=data.get("co_author", ""))
        ws.cell(row=row, column=5, value=data.get("university", ""))
        ws.cell(row=row, column=6, value=data.get("education", ""))
        ws.cell(row=row, column=7, value=data.get("section", ""))
        ws.cell(row=row, column=8, value=data.get("article_title", ""))
        ws.cell(row=row, column=9, value=data.get("supervisor", ""))
        ws.cell(row=row, column=10, value=data.get("file_url", ""))
        ws.cell(row=row, column=11, value=data.get("plagiat_url", ""))
        ws.cell(row=row, column=12, value=data.get("presentation_url", ""))
        
        wb.save(EXCEL_FILE)
        print(f"✅ Заявка #{row-1} добавлена в Excel")
    except Exception as e:
        print(f"Ошибка записи в Excel: {e}")

# ======== КЛАВИАТУРЫ ========
def education_keyboard():
    keyboard = VkKeyboard(one_time=True)
    keyboard.add_button("Бакалавр", color=VkKeyboardColor.PRIMARY)
    keyboard.add_button("Магистр", color=VkKeyboardColor.PRIMARY)
    keyboard.add_line()
    keyboard.add_button("Специалист", color=VkKeyboardColor.PRIMARY)
    keyboard.add_button("Аспирант", color=VkKeyboardColor.PRIMARY)
    return keyboard.get_keyboard()

def section_keyboard():
    keyboard = VkKeyboard(one_time=True)
    keyboard.add_button("Секция 1. Нацбезопасность регионов", color=VkKeyboardColor.PRIMARY)
    keyboard.add_line()
    keyboard.add_button("Секция 2. Финансовая система", color=VkKeyboardColor.PRIMARY)
    keyboard.add_line()
    keyboard.add_button("Секция 3. Управление безопасностью", color=VkKeyboardColor.PRIMARY)
    return keyboard.get_keyboard()

def skip_keyboard():
    """Только для шага с соавтором (шаг 2)"""
    keyboard = VkKeyboard(one_time=True)
    keyboard.add_button("Пропустить", color=VkKeyboardColor.SECONDARY)
    return keyboard.get_keyboard()

def cancel_keyboard():
    """Для обязательных шагов — только кнопка Отмена"""
    keyboard = VkKeyboard(one_time=True)
    keyboard.add_button("Отмена", color=VkKeyboardColor.NEGATIVE)
    return keyboard.get_keyboard()

def main_menu():
    keyboard = VkKeyboard(one_time=False)
    keyboard.add_button("Начать регистрацию", color=VkKeyboardColor.POSITIVE)
    return keyboard.get_keyboard()

def send(user_id, text, keyboard=None):
    try:
        params = {
            "user_id": user_id,
            "message": text,
            "random_id": int(time.time() * 1000)
        }
        if keyboard:
            params["keyboard"] = keyboard
        vk.messages.send(**params)
    except Exception as e:
        print(f"Ошибка отправки: {e}")

# ======== ГЛАВНЫЙ ЦИКЛ ========
print("✅ БОТ ЗАПУЩЕН!")
print(f"📊 Данные сохраняются в файл: {EXCEL_FILE}")
print("⚠️ Руководитель и файлы (шаги 7-10) — ОБЯЗАТЕЛЬНЫ!")

for event in longpoll.listen():
    if event.type == VkBotEventType.MESSAGE_NEW:
        msg = event.object.message
        user_id = msg['from_id']
        text = msg['text'].strip()
        
        if user_id < 0:
            continue
        
        print(f"📨 {user_id}: {text}")
        
        attachments = msg.get('attachments', [])
        
        # ===== КОМАНДА ОТМЕНА =====
        if text == "Отмена":
            if user_id in users:
                del users[user_id]
            send(user_id, "❌ Регистрация отменена. Чтобы начать заново, напишите 'Начать'", main_menu())
            continue
        
        # ===== КОМАНДА "Начать" =====
        if text.lower() in ["начать", "начать регистрацию", "старт"]:
            users[user_id] = {"step": 1}
            send(user_id, "📝 Шаг 1/10: Введите ваше ФИО", cancel_keyboard())
            continue
        
        # ===== ЕСЛИ НЕ В ПРОЦЕССЕ =====
        if user_id not in users:
            send(user_id, "👋 Напишите 'Начать', чтобы зарегистрироваться", main_menu())
            continue
        
        data = users[user_id]
        step = data.get("step", 1)
        
        # ===== ШАГ 1: ФИО =====
        if step == 1:
            data["full_name"] = text
            data["step"] = 2
            send(user_id, "✅ Шаг 2/10: Укажите ФИО соавтора (или нажмите 'Пропустить')", skip_keyboard())
        
        # ===== ШАГ 2: СОАВТОР (можно пропустить) =====
        elif step == 2:
            data["co_author"] = "Нет" if text == "Пропустить" else text
            data["step"] = 3
            send(user_id, "✅ Шаг 3/10: Введите название учебного заведения", cancel_keyboard())
        
        # ===== ШАГ 3: УЧЕБНОЕ ЗАВЕДЕНИЕ =====
        elif step == 3:
            data["university"] = text
            data["step"] = 4
            send(user_id, "✅ Шаг 4/10: Выберите ваше образование:", education_keyboard())
        
        # ===== ШАГ 4: ОБРАЗОВАНИЕ =====
        elif step == 4:
            if text not in ["Бакалавр", "Магистр", "Специалист", "Аспирант"]:
                send(user_id, "⚠️ Пожалуйста, выберите из кнопок!", education_keyboard())
                continue
            data["education"] = text
            data["step"] = 5
            send(user_id, "✅ Шаг 5/10: Выберите секцию конференции:", section_keyboard())
        
        # ===== ШАГ 5: СЕКЦИЯ =====
        elif step == 5:
            short_sections = [
                "Секция 1. Нацбезопасность регионов",
                "Секция 2. Финансовая система",
                "Секция 3. Управление безопасностью"
            ]
            
            full_sections = {
                "Секция 1. Нацбезопасность регионов": 
                    "Секция 1. Пространственное развитие национальной безопасности регионов Российской Федерации: механизмы сбалансированной территориальной политики в условия структурных изменений.",
                "Секция 2. Финансовая система": 
                    "Секция 2. Финансовая система национальной безопасности Российской Федерации в условиях глобальной нестабильности: вызовы и адаптационные механизмы.",
                "Секция 3. Управление безопасностью": 
                    "Секция 3. Теоретические и прикладные аспекты управления безопасностью организаций в условиях цифровой трансформации и экономической неопределенности."
            }
            
            if text not in short_sections:
                send(user_id, "⚠️ Пожалуйста, выберите секцию из кнопок!", section_keyboard())
                continue
            
            data["section"] = full_sections[text]
            data["step"] = 6
            send(user_id, "✅ Шаг 6/10: Введите название статьи", cancel_keyboard())
        
        # ===== ШАГ 6: НАЗВАНИЕ СТАТЬИ =====
        elif step == 6:
            data["article_title"] = text
            data["step"] = 7
            send(user_id, 
                "✅ Шаг 7/10: Укажите ФИО научного руководителя\n\n"
                "⚠️ Этот шаг обязательный!", 
                cancel_keyboard())
        
        # ===== ШАГ 7: РУКОВОДИТЕЛЬ (ОБЯЗАТЕЛЬНО) =====
        elif step == 7:
            # Если пользователь написал "Пропустить" — не пропускаем
            if text == "Пропустить":
                send(user_id, 
                    "⚠️ Научного руководителя нельзя пропустить!\n\n"
                    "👨‍🏫 Введите ФИО научного руководителя:",
                    cancel_keyboard())
                continue
            
            data["supervisor"] = text
            data["step"] = 8
            send(user_id, 
                "📎 Шаг 8/10: Загрузите файл статьи (Word .doc или .docx)\n\n"
                "⚠️ Этот шаг обязательный!\n"
                "Прикрепите файл к сообщению.", 
                cancel_keyboard())
        
        # ===== ШАГ 8: ФАЙЛ СТАТЬИ (WORD) =====
        elif step == 8:
            file_found = False
            for attach in attachments:
                if attach.get('type') == 'doc':
                    doc = attach['doc']
                    file_ext = doc.get('ext', '').lower()
                    
                    if file_ext in ['doc', 'docx']:
                        data["file_url"] = f"https://vk.com/doc{doc['owner_id']}_{doc['id']}"
                        data["step"] = 9
                        file_found = True
                        send(user_id, 
                            "✅ Файл статьи принят!\n\n"
                            "📄 Шаг 9/10: Загрузите справку на антиплагиат (PDF)\n\n"
                            "⚠️ Этот шаг обязательный!\n"
                            "Прикрепите PDF-файл к сообщению.",
                            cancel_keyboard())
                        break
                    else:
                        send(user_id, 
                            f"⚠️ Файл статьи должен быть в формате Word (.doc или .docx)\n"
                            f"Вы прислали: .{file_ext}\n\n"
                            "Попробуйте ещё раз.",
                            cancel_keyboard())
                        file_found = True
                        break
            
            if not file_found:
                send(user_id, 
                    "⚠️ Файл статьи обязателен!\n\n"
                    "📎 Прикрепите файл в формате Word (.doc или .docx)",
                    cancel_keyboard())
        
        # ===== ШАГ 9: СПРАВКА АНТИПЛАГИАТ (PDF) =====
        elif step == 9:
            file_found = False
            for attach in attachments:
                if attach.get('type') == 'doc':
                    doc = attach['doc']
                    file_ext = doc.get('ext', '').lower()
                    
                    if file_ext == 'pdf':
                        data["plagiat_url"] = f"https://vk.com/doc{doc['owner_id']}_{doc['id']}"
                        data["step"] = 10
                        file_found = True
                        send(user_id, 
                            "✅ Справка антиплагиат принята!\n\n"
                            "📊 Шаг 10/10: Загрузите презентацию (PowerPoint .ppt или .pptx)\n\n"
                            "⚠️ Этот шаг обязательный!\n"
                            "Прикрепите файл к сообщению.",
                            cancel_keyboard())
                        break
                    else:
                        send(user_id, 
                            f"⚠️ Справка должна быть в формате PDF\n"
                            f"Вы прислали: .{file_ext}\n\n"
                            "Попробуйте ещё раз.",
                            cancel_keyboard())
                        file_found = True
                        break
            
            if not file_found:
                send(user_id, 
                    "⚠️ Справка на антиплагиат обязательна!\n\n"
                    "📄 Прикрепите PDF-файл",
                    cancel_keyboard())
        
        # ===== ШАГ 10: ПРЕЗЕНТАЦИЯ (POWERPOINT) =====
        elif step == 10:
            file_found = False
            for attach in attachments:
                if attach.get('type') == 'doc':
                    doc = attach['doc']
                    file_ext = doc.get('ext', '').lower()
                    
                    if file_ext in ['ppt', 'pptx']:
                        data["presentation_url"] = f"https://vk.com/doc{doc['owner_id']}_{doc['id']}"
                        file_found = True
                        
                        add_to_excel(data)
                        
                        send(user_id,
                            "✅ РЕГИСТРАЦИЯ УСПЕШНО ЗАВЕРШЕНА!\n\n"
                            f"📌 ФИО: {data['full_name']}\n"
                            f"👥 Соавтор: {data['co_author']}\n"
                            f"🏫 Учебное заведение: {data['university']}\n"
                            f"🎓 Образование: {data['education']}\n"
                            f"📂 Секция: {data['section']}\n"
                            f"📄 Статья: {data['article_title']}\n"
                            f"🧑‍🏫 Руководитель: {data['supervisor']}\n\n"
                            "📎 Все файлы приняты:\n"
                            "✅ Файл статьи (Word)\n"
                            "✅ Справка антиплагиат (PDF)\n"
                            "✅ Презентация (PowerPoint)\n\n"
                            "📝 Ваша заявка сохранена!\n"
                            "Спасибо за участие! 📚",
                            main_menu()
                        )
                        
                        del users[user_id]
                        break
                    else:
                        send(user_id, 
                            f"⚠️ Презентация должна быть в формате PowerPoint (.ppt или .pptx)\n"
                            f"Вы прислали: .{file_ext}\n\n"
                            "Попробуйте ещё раз.",
                            cancel_keyboard())
                        file_found = True
                        break
            
            if not file_found:
                send(user_id, 
                    "⚠️ Презентация обязательна!\n\n"
                    "📊 Прикрепите файл в формате PowerPoint (.ppt или .pptx)",
                    cancel_keyboard())
